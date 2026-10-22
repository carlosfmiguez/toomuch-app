"""Importação em lote: scan de pasta, planilha de mapeamento, aplicação idempotente e tela."""

import io
import tempfile
import zipfile
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse

from . import importer
from .models import Asset, Part, Product


def jpeg_bytes(color=(233, 185, 143)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (40, 50), color).save(buf, format="JPEG")
    return buf.getvalue()


def stl_bytes(extents=(20, 30, 40)) -> bytes:
    import trimesh

    return trimesh.creation.box(extents=extents).export(file_type="stl")


def build_tree(root: Path) -> None:
    (root / "Luminaria Solaris").mkdir()
    (root / "Luminaria Solaris" / "cupula.stl").write_bytes(stl_bytes((180, 180, 200)))
    (root / "Luminaria Solaris" / "base.stl").write_bytes(stl_bytes((80, 80, 60)))
    (root / "Luminaria Solaris" / "solaris-acesa.jpg").write_bytes(jpeg_bytes())
    (root / "Vaso Face").mkdir()
    (root / "Vaso Face" / "vaso-face.stl").write_bytes(stl_bytes())
    (root / "Vaso Face" / "vaso-face-1.jpg").write_bytes(jpeg_bytes((120, 140, 90)))
    (root / "Vaso Face" / "vaso-face-2.jpg").write_bytes(jpeg_bytes((200, 140, 90)))
    (root / "porta-escova.stl").write_bytes(stl_bytes((30, 30, 90)))
    (root / "porta-escova.jpg").write_bytes(jpeg_bytes((90, 90, 90)))
    (root / "Cruz com pedestal").mkdir()
    (root / "Cruz com pedestal" / "cruz.jpg").write_bytes(jpeg_bytes((50, 50, 50)))
    (root / ".DS_Store").write_bytes(b"x")


class ImporterTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_base", verbosity=0)
        cls.user = get_user_model().objects.create_superuser("carlos", "c@example.com", "senha-forte-123")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        build_tree(self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def test_scan_groups_and_suggests(self):
        found = {sp.name: sp for sp in importer.scan(self.root)}
        self.assertEqual(set(found), {"Luminaria Solaris", "Vaso Face", "Porta escova", "Cruz com pedestal"})
        sol = found["Luminaria Solaris"]
        self.assertEqual(sol.category_code, "LUM")
        self.assertEqual(sol.sku_base, "TM-LUM-SOL")
        self.assertEqual(len(sol.models), 2)
        self.assertEqual(len(sol.photos), 1)
        self.assertEqual(found["Vaso Face"].category_code, "VAS")
        self.assertEqual(found["Porta escova"].category_code, "UTI")
        self.assertEqual(found["Cruz com pedestal"].category_code, "REL")
        self.assertEqual(found["Cruz com pedestal"].note, "só fotos, sem arquivo 3D")
        skus = [sp.sku_base for sp in found.values()]
        self.assertEqual(len(skus), len(set(skus)))

    def test_mapping_roundtrip_and_apply_is_idempotent(self):
        products = importer.scan(self.root)
        sheet = importer.write_mapping(products, self.root / "map.xlsx")
        rows = importer.read_mapping(sheet)
        self.assertEqual(len(rows), 4)
        # revisão: renomeia um, pula outro, marca um como licenciado
        from openpyxl import load_workbook

        wb = load_workbook(sheet)
        ws = wb["Mapeamento"]
        for row in ws.iter_rows(min_row=2):
            cells = {c.column_letter: c for c in row}
            if cells["B"].value == "Vaso Face":
                cells["B"].value = "Vaso Face P"
                cells["E"].value = "licensed"
                cells["F"].value = "Designer X, licença comercial"
            if cells["B"].value == "Cruz com pedestal":
                cells["K"].value = "não"
        wb.save(sheet)

        report = importer.apply(self.root, sheet, user=self.user)
        self.assertEqual(report.created_products, 3)
        self.assertEqual(report.skipped_rows, 1)
        self.assertEqual(report.files_added, 8)  # Solaris 3 + Vaso 3 + Porta escova 2
        self.assertEqual(report.errors, [])
        vaso = Product.objects.get(sku_base="TM-VAS-FAC")
        self.assertEqual(vaso.name, "Vaso Face P")
        self.assertEqual(vaso.origin, Product.Origin.LICENSED)
        self.assertEqual(vaso.license_author, "Designer X, licença comercial")
        self.assertEqual(vaso.status, Product.Status.DRAFT)
        self.assertEqual(vaso.variants.count(), 1)
        self.assertEqual(vaso.assets.filter(kind=Asset.Kind.PHOTO).count(), 2)
        self.assertTrue(vaso.assets.filter(kind=Asset.Kind.PHOTO, is_main=True).exists())
        sol = Product.objects.get(sku_base="TM-LUM-SOL")
        self.assertEqual(set(Part.objects.filter(product=sol).values_list("name", flat=True)), {"Cupula", "Base"})
        stl = sol.assets.get(kind=Asset.Kind.STL, name="cupula.stl")
        self.assertEqual(stl.meta["bbox_mm"]["z"], 200.0)
        self.assertTrue(sol.assets.filter(kind=Asset.Kind.PHOTO).first().thumbnail)

        again = importer.apply(self.root, sheet, user=self.user)
        self.assertEqual(again.created_products, 0)
        self.assertEqual(again.reused_products, 3)
        self.assertEqual(again.files_added, 0)
        self.assertEqual(again.files_skipped, 8)

    def test_management_command(self):
        out = io.StringIO()
        call_command("import_folder", str(self.root), mapear=str(self.root / "m.xlsx"), stdout=out)
        self.assertIn("4 produto(s)", out.getvalue())
        self.assertTrue((self.root / "m.xlsx").exists())
        call_command("import_folder", str(self.root), aplicar=str(self.root / "m.xlsx"), usuario="carlos", stdout=out)
        self.assertEqual(Product.objects.filter(status="draft").count(), 4)

    def test_web_flow_with_zip(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            for path in self.root.rglob("*"):
                if path.is_file():
                    zf.write(path, f"modelos/{path.relative_to(self.root)}")
        self.client.login(username="carlos", password="senha-forte-123")
        with override_settings(IMPORT_TMP_DIR=self.root / "work"):
            from . import views_import

            views_import.IMPORT_ROOT = self.root / "work"
            resp = self.client.post(reverse("catalog:import_upload"), {"zip": SimpleUploadedFile("modelos.zip", buf.getvalue())})
            self.assertRedirects(resp, reverse("catalog:import_review"))
            resp = self.client.get(reverse("catalog:import_review"))
            self.assertContains(resp, "Luminaria Solaris")
            self.assertContains(resp, "TM-LUM-SOL")
            post = {"import-0": "on", "name-0": "Luminária Solaris", "cat-0": "LUM", "sku-0": "tm-lum-sol", "origin-0": "own"}
            resp = self.client.post(reverse("catalog:import_review"), post)
            self.assertEqual(resp.status_code, 302)
            self.assertEqual(Product.objects.count(), 1)
            self.assertEqual(Product.objects.get().name, "Luminária Solaris")
            self.assertFalse((self.root / "work").exists() and any((self.root / "work").iterdir()))


class ProductAdminFlowTests(TestCase):
    """O caminho do cadastro pelo admin: peça → 3MF fatiado no próprio produto → variante com composição automática."""

    @classmethod
    def setUpTestData(cls):
        call_command("seed_base", verbosity=0)
        cls.user = get_user_model().objects.create_superuser("carlos", "c@example.com", "senha-forte-123")

    def _inline(self, prefix, rows, total=None):
        data = {f"{prefix}-TOTAL_FORMS": str(total if total is not None else len(rows)), f"{prefix}-INITIAL_FORMS": "0",
                f"{prefix}-MIN_NUM_FORMS": "0", f"{prefix}-MAX_NUM_FORMS": "1000"}
        for i, row in enumerate(rows):
            for k, v in row.items():
                data[f"{prefix}-{i}-{k}"] = v
        return data

    def test_part_file_inline_and_default_bom(self):
        from apps.files3d.tests import make_3mf
        from apps.inventory.models import FilamentSKU

        self.client.login(username="carlos", password="senha-forte-123")
        cat = Product._meta.get_field("category").related_model.objects.get(code="REL")
        fil = FilamentSKU.objects.create(brand="Bambu Lab", manufacturer="Bambu Lab", material="PLA", color_name="Branco")
        base = {
            "name": "Sagrada Família", "sku_base": "TM-REL-SAG", "category": str(cat.pk), "status": "draft",
            "lead_time_days": "5", "origin": "own", "license_cost_per_unit": "0.00",
        }
        # 1) produto com uma peça e uma variante, ainda sem arquivo
        data = {**base, **self._inline("parts", [{"name": "Conjunto", "default_qty": "1", "color_slot": "main"}]),
                **self._inline("assets", []), **self._inline("variants", [{
                    "sku": "TM-REL-SAG-U-BRC", "size": "Único", "main_filament": str(fil.pk), "is_default": "on",
                    "minutes_post": "10", "minutes_assembly": "0", "minutes_packaging": "5", "stock_ready": "0",
                }])}
        # o produto tem dois inlines de Asset com o mesmo prefixo base; o Django numera o segundo como assets-2
        data.update(self._inline("assets-2", []))
        resp = self.client.post(reverse("admin:catalog_product_add"), data)
        self.assertEqual(resp.status_code, 302, resp.content[:2000] if resp.status_code != 302 else "")
        product = Product.objects.get(sku_base="TM-REL-SAG")
        part = product.parts.get()
        variant = product.variants.get()
        self.assertEqual([b.part_id for b in variant.bom_items.all()], [part.pk], "composição automática com a peça")

        # 2) a tela de edição lista a peça no inline de arquivos 3D
        resp = self.client.get(reverse("admin:catalog_product_change", args=[product.pk]))
        self.assertContains(resp, "3MF fatiado por peça")
        self.assertContains(resp, "abrir cadastro completo")

        # 3) sobe o 3MF fatiado pelo inline do produto, escolhendo a peça
        data = {**base, **self._inline("parts", []), **self._inline("variants", []), **self._inline("assets-2", [])}
        data["parts-INITIAL_FORMS"] = "0"
        data.update(self._inline("assets", [{"part": str(part.pk), "kind": "3mf_sliced", "is_production": "on", "version": "1"}]))
        data["assets-0-file"] = SimpleUploadedFile("conjunto_A1.gcode.3mf", make_3mf())
        resp = self.client.post(reverse("admin:catalog_product_change", args=[product.pk]), data)
        self.assertEqual(resp.status_code, 302)
        asset = part.production_assets().get()
        self.assertEqual(asset.product_id, product.pk)
        self.assertIsNotNone(asset.grams_total)
        self.assertTrue(product.catalog_checklist()[0][1], "3MF fatiado em todas as peças")
